"""Import machine profiles from a "machine speeds" commissioning workbook.

The workbook has one sheet per machine. Each row is a label in the first
column, one value per operating point in the columns after it and a unit
last; a row without a label gives the previous row's value in Hz:

    Machine Type                 <name>         <name>
    1. Operating parameters
    Bowl speed                   3000           2800      rpm
    Differential speed           8              8         rpm
    2. Gearbox and belt data
    Gearbox Type                 Cyclo - 2 stages …
    Cyclo 1st stage ratio, K1    -60            -60       ratio
    …
    3. Exciting speed
    Main motor speed             1445.9         …         rpm
                                 24.1           …         Hz

The gearbox type (and a second motor's belt) picks a template
(physics/templates.py), whose formulas the sheet's parameters fill in. The
speeds of section 3 become each operating point's reference values, so the
profile editor shows where the formulas and the sheet agree. Rows that are
not recognised are reported, not guessed at.
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from ..physics.profile import MachineProfileData
from ..physics.templates import TEMPLATES


class WorkbookError(ValueError):
    """Not a workbook, or one without any sheet in the expected layout."""


#: Limits on what a workbook may unpack to: an .xlsx is a zip, whose small
#: download can expand to gigabytes of XML. Commissioning sheets unpack to
#: well under a megabyte.
MAX_UNPACKED_BYTES = 50 * 1024 * 1024
MAX_ENTRIES = 2000


def _check_archive(content: bytes) -> None:
    """Refuse an archive that would unpack to more than the limits allow,
    going by the sizes its directory declares (zipfile itself refuses
    entries that inflate beyond their declared size)."""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
    except (zipfile.BadZipFile, OSError, ValueError) as exc:
        raise WorkbookError("not an .xlsx workbook") from exc
    if len(entries) > MAX_ENTRIES or sum(e.file_size for e in entries) > MAX_UNPACKED_BYTES:
        raise WorkbookError("the workbook unpacks to more than a commissioning sheet can hold")


@dataclass
class ImportedProfile:
    profile: MachineProfileData
    sheet: str
    warnings: list[str] = field(default_factory=list)


@dataclass
class _Row:
    label: str
    values: list[float]
    unit: str
    text: list[str]


def _norm(label: str) -> str:
    """Lower case, Ø spelled out, punctuation and spaces collapsed."""
    text = label.lower().replace("ø", " diameter ").replace("Ø", " diameter ")
    return re.sub(r"[^a-z0-9.]+", " ", text).strip()


def _rows(ws: Any) -> list[_Row]:
    out: list[_Row] = []
    for raw in ws.iter_rows(values_only=True):
        cells = list(raw)
        if not any(c is not None and str(c).strip() for c in cells):
            continue
        label = str(cells[0]).strip() if cells and cells[0] is not None else ""
        values: list[float] = []
        text: list[str] = []
        unit = ""
        for cell in cells[1:]:
            if isinstance(cell, bool) or cell is None:
                continue
            if isinstance(cell, int | float):
                if unit:  # numbers after the unit column are notes, not values
                    break
                # Six decimals keep a sheet's precision and drop the noise of
                # its computed cells (1464.5166959578207).
                values.append(round(float(cell), 6))
            elif str(cell).strip():
                if values:
                    unit = str(cell).strip()
                else:
                    text.append(str(cell).strip())
        out.append(_Row(label, values, unit, text))
    return out


# Section 2 parameters: (key, test on the normalised label).
_GEARBOX_KEYS = [
    ("K1", lambda s: "k1" in s and "k1.k2" not in s),
    ("K2", lambda s: "k2" in s and "k1.k2" not in s),
]
# Section 3 rows → component keys (planetary: the "Vi" row is the pinion).
_REFERENCE_KEYS = [
    ("2nd motor belt flexing", "belt2_flex"),
    ("2nd motor belt frequency", "belt2"),
    ("belt flexing", "belt_flex"),
    ("belt frequency", "belt"),
    ("main motor speed", "main_motor"),
    ("secondary motor speed", "secondary_motor"),
    ("intermediate shaft speed", "intermediate_shaft"),
    ("pinion speed", "pinion"),
    ("conveyor speed", "scroll"),
    ("bowl speed", "bowl"),
]


def _sheet_profile(title: str, rows: list[_Row]) -> ImportedProfile | None:
    warnings: list[str] = []
    name = ""
    gearbox = ""
    section = 0
    second_motor = False
    params: dict[str, list[float]] = {}
    references: dict[str, list[float]] = {}
    bowl: list[float] | None = None
    diff: list[float] | None = None
    product: list[float] | None = None

    for row in rows:
        s = _norm(row.label)
        if not s:
            continue  # the Hz line under a speed: implied by the rpm line
        if re.match(r"^1\b.*operating", s):
            section = 1
            continue
        if re.match(r"^2\b.*(gear|belt)", s):
            section = 2
            continue
        if re.match(r"^3\b.*(speed|frequenc)", s):
            section = 3
            continue
        if s.startswith("machine type"):
            name = (row.text or [title])[0]
            continue
        if s.startswith("gearbox type"):
            gearbox = " ".join(row.text).lower()
            continue
        if not row.values:
            warnings.append(f"'{row.label}': no value, ignored")
            continue

        if section in (0, 1) and s.startswith("bowl speed"):
            bowl = row.values
        elif section in (0, 1) and s.startswith("differential speed"):
            diff = row.values
        elif section == 2:
            if s.startswith("2nd motor"):
                second_motor = True
            prefix = "2" if s.startswith("2nd motor") else ""
            if "k1.k2" in s:
                product = row.values
            elif key := next((k for k, test in _GEARBOX_KEYS if test(s)), None):
                params[key] = row.values
            elif "pulley" in s and "motor" in s.removeprefix("2nd motor"):
                params[f"D{prefix}_motor"] = row.values
            elif "pulley" in s and "machine" in s:
                params[f"D{prefix}_machine"] = row.values
            elif "belt length" in s:
                # The second belt's length follows the second motor's pulleys.
                params["L2_belt" if second_motor else "L_belt"] = row.values
            elif "pulley ratio" in s:
                pass  # derived from the diameters
            else:
                warnings.append(f"'{row.label}': not recognised, ignored")
        elif section == 3:
            key = next((k for text, k in _REFERENCE_KEYS if text in s), None)
            if key is None:
                warnings.append(f"'{row.label}': not recognised, ignored")
            elif row.unit.lower() in ("rpm", "1/min", "u/min", ""):
                references[key] = row.values
        else:
            warnings.append(f"'{row.label}': not recognised, ignored")

    if bowl is None:
        return None  # not a machine-speeds sheet

    if second_motor or "planet" in gearbox:
        template_id = "decanter-planetary"
    else:
        template_id = "decanter-cyclo"
        if "cyclo" not in gearbox:
            warnings.append(
                f"gearbox type '{gearbox or '–'}' is not known; the Cyclo template was used"
            )
    doc: dict[str, Any] = {k: v for k, v in TEMPLATES[template_id].items()}
    if "pinion" in references and template_id == "decanter-cyclo":
        references["intermediate_shaft"] = references.pop("pinion")
    if "intermediate_shaft" in references and template_id == "decanter-planetary":
        references["pinion"] = references.pop("intermediate_shaft")

    n_points = len(bowl)
    if diff is not None:
        params["d"] = diff
    if "K1" in params and "K2" in params and product is not None:
        expected = params["K1"][0] * params["K2"][0]
        if abs(expected - product[0]) > 1e-6 * max(1.0, abs(expected)):
            warnings.append(f"K1.K2 is {product[0]:g} on the sheet, but K1 × K2 = {expected:g}")

    parameters = []
    point_values: list[dict[str, float]] = [{} for _ in range(n_points)]
    template_keys = {p["key"] for p in doc["parameters"]}
    for p in doc["parameters"]:
        values = params.get(p["key"])
        if values is None:
            if p["key"] != "f_mains":
                warnings.append(
                    f"{p['label']} ({p['key']}) is not on the sheet; example value kept"
                )
            parameters.append(dict(p))
            continue
        parameters.append({**p, "value": values[0]})
        for i, v in enumerate(values[1:n_points], start=1):
            if v != values[0]:
                point_values[i][p["key"]] = v
    for key in sorted(set(params) - template_keys):
        warnings.append(f"'{key}' does not fit the {template_id} template; ignored")

    component_keys = {c["key"] for c in doc["components"]}
    points = []
    for i, rpm in enumerate(bowl):
        refs = {k: v[i] for k, v in references.items() if k in component_keys and i < len(v)}
        points.append(
            {
                "label": f"{rpm:g} rpm",
                "bowl_rpm": rpm,
                "parameters": point_values[i],
                "reference_rpm": refs,
            }
        )

    machine = (name or title).strip()
    doc.update(
        name=machine[:80],
        description=f"Imported from the sheet '{title}' ({doc['name']} template).",
        parameters=parameters,
        operating_points=points,
        match_patterns=[machine.lower()[:60]] if len(machine) >= 2 else [],
    )
    return ImportedProfile(MachineProfileData.model_validate(doc), title, warnings)


def read_machine_speeds(content: bytes) -> list[ImportedProfile]:
    """Every machine sheet of the workbook; WorkbookError if there is none."""
    _check_archive(content)
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
        raise WorkbookError("not an .xlsx workbook") from exc
    try:
        out = [p for ws in wb.worksheets if (p := _sheet_profile(ws.title, _rows(ws))) is not None]
    finally:
        wb.close()
    if not out:
        raise WorkbookError("no sheet in the machine speeds layout (no 'Bowl speed' row)")
    return out
